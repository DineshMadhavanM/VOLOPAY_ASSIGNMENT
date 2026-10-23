from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

# Initialize presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette (Dark Professional Fintech Theme)
BG_COLOR = RGBColor(11, 15, 25)       # #0b0f19
CARD_COLOR = RGBColor(17, 24, 39)     # #111827
TEXT_WHITE = RGBColor(248, 250, 252)  # #f8fafc
TEXT_MUTED = RGBColor(148, 163, 184)  # #94a3b8
ACCENT_INDIGO = RGBColor(99, 102, 241) # #6366f1
ACCENT_CYAN = RGBColor(6, 182, 212)    # #06b6d4
ACCENT_EMERALD = RGBColor(16, 185, 129) # #10b981
ACCENT_AMBER = RGBColor(245, 158, 11)   # #f59e0b
ACCENT_ROSE = RGBColor(239, 68, 68)     # #ef4444

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def add_header(slide, badge_text, title_text):
    # Badge
    tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
    tf = tx_box.text_frame
    p = tf.paragraphs[0]
    p.text = badge_text.upper()
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    # Title
    tx_box2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.7))
    tf2 = tx_box2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE

def add_footer(slide, current_num):
    tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4))
    tf = tx_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"SpendNest Revenue Intelligence Inbox — Volopay Assessment | Slide {current_num} of 5"
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_MUTED

# Slide 1
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide1)
add_header(slide1, "Slide 1 • Problem Definition & Ground-Truth Evidence", "1. The Manager Problem Solved & 3 Key Findings")
add_footer(slide1, 1)

content_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.0))
tf = content_box.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "The Problem: Sales managers receive messy telephony & calendar logs where labels contradict reality, background dials inflate activity counts, and vendor meetings are mistakenly counted as pipeline."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_MUTED
p.space_after = Pt(14)

p = tf.add_paragraph()
p.text = "Three Critical Findings Linked to Exact Source Evidence:"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_INDIGO
p.space_after = Pt(10)

findings = [
    ("Finding 1: Vendor Evaluation Mistagged as Sales Pipeline (M02 - LedgerBridge)",
     "Evidence (Turn 1): 'I work on strategy and partnerships at SpendNest... We want to understand whether LedgerBridge can give us compliant invoicing coverage in Europe...'\nImpact: SpendNest is the buyer! Upstream CRM tagged this as sales pipeline. Our system routes it to Vendor Evaluation / Internal, protecting quota and forecasting accuracy."),
    ("Finding 2: Telephony User Identity Mismatch (Call C06)",
     "Evidence (Verbatim quote): 'Good evening, Mr Nikhil... Ravi here. I'll say at the start that this is a cold call.'\nImpact: Telephony metadata lists rep as Meera, but spoken audio is Ravi. Flagged as Rep Coaching / Data Discrepancy to prevent commission and coaching misallocation."),
    ("Finding 3: Product Capability Disconnect (Meeting M01 - Harborpath Advisory)",
     "Evidence (Turn 1 & 2): Partner seeks remote bank account opening. SpendNest clarifies: 'we provide virtual business accounts for operational spend, not licensed retail bank accounts...'\nImpact: Channel partner risk. Must deliver compliance positioning before launching partner referrals.")
]

for title, desc in findings:
    p = tf.add_paragraph()
    p.text = "• " + title
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    
    p2 = tf.add_paragraph()
    p2.text = "   " + desc
    p2.font.size = Pt(11)
    p2.font.color.rgb = TEXT_MUTED
    p2.space_after = Pt(8)

# Slide 2
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide2)
add_header(slide2, "Slide 2 • UX Design Philosophy & System Boundaries", "2. Interface Choices & Deliberate Omissions")
add_footer(slide2, 2)

content_box = slide2.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.0))
tf = content_box.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Manager-Centric UX Design: Prioritize high-signal exceptions, explicit denominators, and evidence drill-down over vanity metric charts."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_MUTED
p.space_after = Pt(14)

p = tf.add_paragraph()
p.text = "Key Interface Decisions:"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_CYAN
p.space_after = Pt(8)

choices = [
    ("Built-In Replay Simulator Toolbar", "5 interactive buttons to test idempotent repeat imports, manual route overrides, and late-arriving update batches directly in the UI."),
    ("Strict Mathematical Denominators", "Separates Call Reps (SDRs/AEs dialing) from Meeting Organizers (calendar coordinators). Transcript coverage explicitly measured over all 44 selected records."),
    ("Deep Evidence Drawer with Formatted Turns", "Allows inspecting verbatim turns, elapsed timecodes, quotes, limitations, and assigned action owners behind every recommendation.")
]

for title, desc in choices:
    p = tf.add_paragraph()
    p.text = "• " + title + ": "
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.text += desc
    p.space_after = Pt(8)

p = tf.add_paragraph()
p.text = "Deliberate Omissions (Design Discipline):"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_AMBER
p.space_before = Pt(10)
p.space_after = Pt(8)

omissions = [
    ("Zero Fabricated Summaries", "For the 26 records with null transcripts or empty turns, we deliberately declined to generate fake summaries. They are logged transparently as Intake / No Audio Supplied."),
    ("No Standalone Date Filter", "All supplied interactions fall on 21 August 2026 (Asia/Kolkata). Providing a date picker adds zero analytical value and was omitted.")
]

for title, desc in omissions:
    p = tf.add_paragraph()
    p.text = "• " + title + ": "
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.text += desc
    p.space_after = Pt(6)

# Slide 3
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide3)
add_header(slide3, "Slide 3 • Routing Architecture & Difficult Decision", "3. Routing Taxonomy & The Difficult Case (M07)")
add_footer(slide3, 3)

content_box = slide3.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.0))
tf = content_box.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Action-Oriented Routing Taxonomy (6 Distinct Operational Destinations):"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_EMERALD
p.space_after = Pt(10)

routes = [
    ("1. Deal Next Steps / Pipeline", "Active commercial evaluations with stated pricing/budget (M03 Asterloom Systems, C03 Milan API, C08 Raghav referral)."),
    ("2. Customer Success / Escalations", "Urgent customer bug escalations and account migrations (M07 Willowfield Trust, M04 Cedarwave Distribution)."),
    ("3. Partnerships / Channel Advisory", "Referral partners incorporating overseas businesses (C07 Willowbank Advisory, M01 Harborpath Advisory)."),
    ("4. Vendor & Internal Enablement", "Procurement reviews (M02 LedgerBridge) & internal staff training (M10 Billing walkthrough)."),
    ("5. Low Priority / Voicemail / Gatekeeper", "Gatekeeper transfers and voicemails (C01, C02, C04, C06, C10)."),
    ("6. Intake / Unreviewed", "Short dials & empty meetings awaiting audio (B01-B20, M05, M08, M11-M14).")
]

for title, desc in routes:
    p = tf.add_paragraph()
    p.text = title + " — "
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.text += desc
    p.space_after = Pt(4)

p = tf.add_paragraph()
p.text = "The Difficult Routing Case: Willowfield Trust (M07)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_ROSE
p.space_before = Pt(12)
p.space_after = Pt(6)

p = tf.add_paragraph()
p.text = "• In Delivery Batch 1, M07 arrived with null turns and was provisionally labeled as an unreviewed routine check-in.\n• In Delivery Batch 2 (Rev 2), 27 turns arrived showing Customer A screen-sharing critical software defects: 1) payment description edits reverting upon save, and 2) document upload failing when recreating payments.\n• Decision: Shifted immediately from Routine Check-in to Customer Success / Bug Escalation. Under DATA_CONTRACT.md, manager overrides survive batch replays and browser reloads."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_MUTED

# Slide 4
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide4)
add_header(slide4, "Slide 4 • Audit Trail & Replay Verification", "4. Audit Verification of Key Numbers & Claims")
add_footer(slide4, 4)

content_box = slide4.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.0))
tf = content_box.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "Mathematical Verification Against 01_initial.json and 02_update.json:"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_CYAN
p.space_after = Pt(12)

metrics = [
    ("Unique Source Records", "44 records (Rev 1) -> 44 records (Rev 2)", "Keyed strictly by (source_type, source_id). Higher revision updates in place."),
    ("Transcript Coverage", "17 / 44 (38.6%) -> 18 / 44 (40.9%)", "M07 arrives in Rev 2 with 27 turns, raising coverage from 17 to 18."),
    ("Calls vs. Meetings Split", "30 Calls & 14 Meetings", "Partitioned strictly by source_type; remains constant across batches."),
    ("Source Connection Rate", "9 / 30 calls (30.0%)", "Calls with is_connected=true / 30 calls. Exactly 8 had multi-turn prospect dialogue.")
]

for title, val, formula in metrics:
    p = tf.add_paragraph()
    p.text = "• " + title + ": "
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.text += f"{val} ({formula})"
    p.space_after = Pt(8)

p = tf.add_paragraph()
p.text = "Repeat Import Idempotence Test:"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_EMERALD
p.space_before = Pt(12)
p.space_after = Pt(6)

p = tf.add_paragraph()
p.text = "Re-importing Batch 1 with altered record order confirms strict idempotence: 0 duplicates added, 44 duplicates safely rejected, zero metric inflation."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MUTED

# Slide 5
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_background(slide5)
add_header(slide5, "Slide 5 • Engineering Ethics & AI Disclosure", "5. Trade-offs, Rejected Decision & Next Steps")
add_footer(slide5, 5)

content_box = slide5.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.0))
tf = content_box.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "The Deliberately Rejected Decision: Auto-Estimating Deal Pipeline ARR"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_ROSE
p.space_after = Pt(8)

p = tf.add_paragraph()
p.text = "• Alternative Considered: Using an LLM to hallucinate contract values based on prospect employee counts (e.g. estimating Asterloom at $45,000 ARR).\n• Why Rejected: The prompt strictly prohibits inventing pipeline value or deal outcomes. Speculative revenue figures destroy executive trust during audits.\n• Accepted Trade-off: We report qualitative commercial intent and verbatim financial statements (e.g. Asterloom's INR 26-39 lakh monthly vendor spend) without guessing ARR."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_MUTED
p.space_after = Pt(14)

p = tf.add_paragraph()
p.text = "AI Implementation & Architecture Disclosure:"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_INDIGO
p.space_after = Pt(8)

p = tf.add_paragraph()
p.text = "• Pre-generated vs. Dynamic: Conversation analyses and evidence citations were prepared beforehand. Batch ingestion, revision deduplication, manual override persistence, and ZIP manifest generation run completely dynamically client-side in the browser.\n• Approximate Time Spent: ~7 hours total (modeling, UI design, contract compliance, documentation).\n• Known Limitations & Future Scope: Cross-call speaker identity is not unified upstream. Future release will integrate live webhooks for automated recording synchronization."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_MUTED

# Save presentation
output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'slides', 'presentation.pptx')
prs.save(output_path)
print(f"PowerPoint Presentation successfully created at: {output_path}")
